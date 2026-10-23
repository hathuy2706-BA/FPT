# -*- coding: utf-8 -*-
"""
Cập nhật bảng Kịch Bản Chuẩn - Tech Spec sang chuẩn nghiệp vụ MUA KÈM PHỤ KIỆN (ATTACHED ADD-ON)
Lưu ý nghiệp vụ cốt lõi:
- ĐÂY KHÔNG PHẢI LÀ QUÀ TẶNG, MÀ LÀ BÁN KÈM PHỤ KIỆN VỚI GIÁ MUA KÈM ƯU ĐÃI KHI MUA CAMERA CHÍNH
- Giá mua kèm ưu đãi hiện tại đang áp dụng là 0đ (hoặc phụ thu 50k cho bản Premium), 
  nhưng trong tương lai có thể cấu hình linh hoạt thay đổi (VD: 29k, 49k, 99k,...)
- Hạn mức mua kèm ưu đãi (Add-on Quota) gắn liền với số lượng Camera chính trong giỏ.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import csv
import os

modules_biz = [
    {
        "module": "1. Tổ hợp số lượng mua kèm (Tỷ lệ 1:1 chuẩn)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Mua đúng tỷ lệ 1:1: Mua 1 Camera Play 4 và chọn mua kèm 1 Charm tiêu chuẩn (Basic)",
                "pre": "Giỏ hàng có 1 Camera Play 4 (giá niêm yết). Kho còn Charm tiêu chuẩn.",
                "rule": "1 Camera Play 4 kích hoạt 01 suất mua kèm 01 Charm tiêu chuẩn với mức giá mua kèm ưu đãi (hiện tại áp dụng 0đ, có thể cấu hình linh hoạt). Hạn mức mua kèm giá ưu đãi = Số lượng Camera trong giỏ hàng (1 Cam = 1 suất mua kèm ưu đãi).",
                "alt": "Nếu mẫu Charm tiêu chuẩn khách chọn bị hết hàng: Cho phép khách chọn đổi sang mẫu Charm tiêu chuẩn khác còn hàng trong hạn mức mua kèm.",
                "biz_question": "Đội Kinh doanh dự kiến trong tương lai mức giá mua kèm ưu đãi này sẽ giữ 0đ hay chuyển sang thu phí ưu đãi (ví dụ: mua kèm với giá 29k, 49k)?"
            },
            {
                "br": "BR-02",
                "scenario": "Vượt hạn mức mua kèm (Cùng mẫu): Mua 1 Cam, khách chọn 2 Charm cùng Mẫu A",
                "pre": "Đã có 1 Camera và 1 Charm A (giá mua kèm ưu đãi 0đ) trong giỏ. Khách bấm chọn thêm 1 Charm A nữa.",
                "rule": "Tách 2 Charm A thành 2 dòng riêng: 01 chiếc tính theo giá mua kèm ưu đãi (hiện tại 0đ, theo hạn mức Camera), 01 chiếc tính giá bán lẻ gốc niêm yết 150.000đ. Không tính giá trung bình cộng.",
                "alt": "Nếu khách bấm giảm bớt 1 Charm A: Hệ thống luôn ưu tiên giữ lại chiếc hưởng giá mua kèm ưu đãi cho khách, bỏ chiếc tính giá bán lẻ.",
                "biz_question": "Khi khách chọn vượt quá số lượng mua kèm ưu đãi, Kinh doanh muốn hệ thống tự động tính giá bán lẻ chiếc thứ 2 hay hiển thị cảnh báo giới hạn?"
            },
            {
                "br": "BR-03",
                "scenario": "Vượt hạn mức mua kèm (Khác mẫu - Khác giá): Mua 1 Cam, chọn 1 Charm A (100k) và 1 Charm B (150k)",
                "pre": "1 Camera Play 4 trong giỏ, khách tick chọn 1 Charm A (giá bán lẻ 100k) và 1 Charm B (giá bán lẻ 150k).",
                "rule": "Ưu tiên quyền lợi tốt nhất cho khách (Best-deal): Tự động áp dụng giá mua kèm ưu đãi (0đ) cho Charm có giá trị cao nhất (Charm B - 150k), Charm còn lại (Charm A - 100k) tính theo giá bán lẻ gốc.",
                "alt": "Nếu khách bỏ chọn Charm B: Hệ thống tự động chuyển Charm A sang hưởng suất mua kèm ưu đãi ngay lập tức.",
                "biz_question": "Kinh doanh muốn ưu tiên áp dụng giá mua kèm ưu đãi cho món có giá trị cao nhất (Best-deal) hay món có giá trị thấp nhất để tối ưu doanh thu phụ kiện?"
            },
            {
                "br": "BR-04",
                "scenario": "Dưới hạn mức mua kèm: Mua 3 Camera nhưng khách chỉ chọn mua kèm 1 Charm Mẫu A",
                "pre": "Giỏ hàng có 3 Camera (được 3 suất mua kèm ưu đãi), nhưng khách chỉ chọn mua kèm 1 Charm A.",
                "rule": "Không ép buộc mua đủ số lượng mua kèm. Khách có toàn quyền mua ít hơn hạn mức. Hệ thống hiển thị lời nhắc gợi ý mua kèm nhưng không chặn luồng thanh toán.",
                "alt": "Khách bỏ qua lời nhắc và thanh toán: Đơn hàng chỉ xuất 3 Camera + 1 Charm. Quyền mua kèm ưu đãi của 2 Charm còn lại sẽ hết hiệu lực khi đơn hàng hoàn tất.",
                "biz_question": "Có cần gửi thông báo nhắc khách vẫn còn suất mua phụ kiện giá ưu đãi không, hay để khách tự quyết định không mua kèm?"
            },
            {
                "br": "BR-05",
                "scenario": "Không mua kèm phụ kiện: Mua 2 Camera nhưng không chọn bất kỳ Charm nào",
                "pre": "Có 2 Camera Play 4 trong giỏ, khách không chọn phụ kiện mua kèm nào và bấm đi đến thanh toán.",
                "rule": "Chính sách mua kèm là ưu đãi 'Mua Kèm Ngay Cùng Đơn' (Instant Add-on), KHÔNG quy đổi thành tiền mặt, KHÔNG tự động bảo lưu thành voucher mua phụ kiện ở đơn hàng sau.",
                "alt": "Tại bước thanh toán: Hiển thị bảng xác nhận nhanh: 'Bạn đang không chọn phụ kiện mua kèm giá ưu đãi. Bạn có muốn chọn thêm phụ kiện trước khi thanh toán?'.",
                "biz_question": "Nếu sau này khách nhận Camera rồi mới muốn mua bổ sung phụ kiện theo giá mua kèm ưu đãi thì CSKH có hỗ trợ tạo đơn bổ sung không?"
            },
            {
                "br": "BR-06",
                "scenario": "Giảm số lượng Camera (Đồng bộ ngược): Đang có 2 Cam + 2 Charm A (giá ưu đãi), giảm Cam xuống 1",
                "pre": "Giỏ hàng có 2 Camera + 2 Charm A (cả 2 đều giá mua kèm ưu đãi 0đ). Khách bấm giảm bớt 1 Camera.",
                "rule": "Hạn mức mua kèm ưu đãi giảm từ 2 xuống 1. Giữ nguyên 2 Charm A trong giỏ nhưng chiếc thứ 2 tự động chuyển thành tính giá bán lẻ niêm yết (150.000đ). Không tự ý xóa sản phẩm của khách.",
                "alt": "Khách có quyền bấm xóa dòng Charm tính giá bán lẻ 150.000đ nếu không có nhu cầu mua thêm ngoài hạn mức.",
                "biz_question": "Kinh doanh muốn giữ lại phụ kiện và chuyển sang giá bán lẻ, hay tự động giảm số lượng phụ kiện mua kèm tương ứng với số Camera?"
            },
            {
                "br": "BR-07",
                "scenario": "Giảm Camera khi đang chọn nhiều mẫu phụ kiện khác nhau: Đang có 2 Cam + 1 Charm A + 1 Charm B, giảm Cam xuống 1",
                "pre": "Giỏ có 2 Cam + 1 Charm A + 1 Charm B (cả 2 đều giá ưu đãi 0đ). Khách giảm Camera xuống 1.",
                "rule": "Hạn mức ưu đãi chỉ còn 1: Giữ món có giá bán lẻ cao nhất hưởng giá mua kèm ưu đãi (0đ), món còn lại chuyển về giá bán lẻ gốc. Nếu 2 món bằng giá, ưu tiên món khách đã chọn trước.",
                "alt": "Cho phép khách nhấp chuột trực tiếp để hoán đổi quyền hưởng giá ưu đãi giữa Mẫu A và Mẫu B một cách dễ dàng.",
                "biz_question": "Khi giảm số lượng Camera, Kinh doanh muốn hệ thống tự động gán ưu đãi theo thuật toán hay bật bảng hỏi để khách tự chọn mẫu được hưởng ưu đãi?"
            },
            {
                "br": "BR-08",
                "scenario": "Mua số lượng lớn: Khách mua 10 Camera, muốn chia phụ kiện mua kèm cho cả 3 mẫu Charm A, B, C",
                "pre": "Khách mua số lượng lớn (10 Camera Play 4, được 10 suất mua kèm phụ kiện giá ưu đãi).",
                "rule": "Cho phép khách nhập trực tiếp số lượng từng mẫu phụ kiện mua kèm vào ô số lượng. Tổng số lượng Charm hưởng giá mua kèm ưu đãi tối đa không vượt quá 10 chiếc.",
                "alt": "Hiển thị thanh tiến trình trực quan: 'Đã phân bổ: 7/10 phụ kiện mua kèm ưu đãi'. Nếu nhập vượt quá 10, số lượng dư ra sẽ tự động tính theo giá bán lẻ.",
                "biz_question": "Chính sách mua kèm số lượng lớn (trên 10 thiết bị) có áp dụng chung chính sách mua kèm này hay áp dụng theo báo giá đại lý/B2B riêng?"
            },
            {
                "br": "BR-09",
                "scenario": "Đổi mẫu phụ kiện mua kèm: Đang chọn 3 Cam + 3 Charm (1A, 1B, 1C), muốn đổi C thành A",
                "pre": "Đã chọn tối đa 3 Charm mua kèm (1A, 1B, 1C). Hạn mức mua kèm ưu đãi khả dụng = 0.",
                "rule": "Khi khách giảm số lượng Mẫu C về 0, hạn mức mua kèm ưu đãi lập tức được giải phóng. Khách bấm tăng Mẫu A lên 2 chiếc thì chiếc thứ 2 được áp dụng giá mua kèm ưu đãi.",
                "alt": "Nếu khách tăng Mẫu A trước khi giảm Mẫu C: Chiếc A thứ 2 tạm tính giá bán lẻ 150k cho đến khi khách giảm bớt Mẫu C.",
                "biz_question": "Giao diện đổi mẫu phụ kiện mua kèm có cần hiển thị nút 'Đổi mẫu mua kèm' trực tiếp tại từng dòng sản phẩm không?"
            },
            {
                "br": "BR-10",
                "scenario": "Xóa Camera khỏi giỏ: Đang có 1 Cam + 1 Charm mua kèm, khách bấm xóa Camera",
                "pre": "Có 1 Camera Play 4 + 1 Charm mua kèm giá ưu đãi trong giỏ. Khách bấm nút 'Xóa' Camera.",
                "rule": "Quy tắc ràng buộc chặt (Strict Binding): Phụ kiện mua kèm giá ưu đãi không thể tồn tại độc lập khi không có sản phẩm chính. Xóa Camera đồng nghĩa mất điều kiện mua kèm giá ưu đãi.",
                "alt": "Hệ thống hỏi khách: 'Xóa Camera sẽ mất quyền mua phụ kiện giá ưu đãi. Bạn có muốn giữ lại phụ kiện với giá bán lẻ 150.000đ không? [Xóa tất cả] [Giữ Charm giá bán lẻ]'.",
                "biz_question": "Kinh doanh muốn tự động xóa toàn bộ phụ kiện mua kèm khi xóa Camera chính, hay giữ lại phụ kiện và chuyển sang giá bán lẻ?"
            }
        ]
    },
    {
        "module": "2. Phụ thu nâng cấp phụ kiện (Charm Premium +50k)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Mua kèm kết hợp Bản tiêu chuẩn & Bản nâng cấp: Mua 2 Cam, chọn 1 Charm Basic + 1 Charm Premium",
                "pre": "2 Camera Play 4 (được 2 suất mua kèm ưu đãi). Kho còn cả mẫu Basic và mẫu Premium.",
                "rule": "Mỗi Camera cho phép mua kèm 1 Charm tiêu chuẩn với giá ưu đãi (hiện tại 0đ) HOẶC nâng cấp lên Charm Premium với mức giá phụ thu ưu đãi +50.000đ/chiếc (thay vì giá bán lẻ gốc 200.000đ). Mức phụ thu này có thể cấu hình linh hoạt.",
                "alt": "Bảng kê đơn hàng hiển thị rõ: Tiền Camera + Phụ thu nâng cấp Charm Premium 50.000đ (Tiết kiệm 150.000đ so với giá bán lẻ).",
                "biz_question": "Khoản thu 50.000đ này kế toán sẽ ghi nhận vào doanh thu bán lẻ phụ kiện hay phân bổ tính vào gói sản phẩm Camera?"
            },
            {
                "br": "BR-02",
                "scenario": "Vượt hạn mức khi chọn cả 2 loại: Mua 1 Cam, chọn cả 1 Charm Basic + 1 Charm Premium",
                "pre": "1 Camera Play 4 (chỉ có 1 suất mua kèm ưu đãi). Khách chọn 2 phụ kiện khác phân khúc.",
                "rule": "Áp dụng suất mua kèm ưu đãi cho món mang lại quyền lợi cao nhất: Charm Basic giá ưu đãi 0đ + Charm Premium tính giá bán lẻ gốc 200k (Tổng tiền phụ kiện = 200k).",
                "alt": "Nếu khách muốn áp dụng suất ưu đãi cho Premium (phụ thu 50k) và trả tiền gốc cho Basic (150k): Tổng tiền phụ kiện vẫn là 200k.",
                "biz_question": "Kinh doanh có đồng ý cho khách chủ động chọn món nào được hưởng suất mua kèm ưu đãi (Basic hay Premium) trong trường hợp này không?"
            },
            {
                "br": "BR-03",
                "scenario": "Đổi ý nâng cấp (Swap): Đang chọn Charm Basic, khách bấm đổi sang Charm Premium",
                "pre": "Giỏ có 1 Cam + 1 Charm Basic giá ưu đãi 0đ. Khách tick chọn sang mẫu Premium.",
                "rule": "Hệ thống tự động hủy chọn mẫu Basic và kích hoạt mức giá mua kèm ưu đãi cho Premium (phụ thu +50.000đ). Tổng tiền thanh toán tăng thêm đúng 50.000đ.",
                "alt": "Giao diện cập nhật tức thì kèm thông báo: 'Đã chuyển sang mua kèm phiên bản phụ kiện Premium (+50.000đ)'.",
                "biz_question": "Có cần thiết kế nhãn 'Nâng cấp chỉ +50k' nổi bật để thúc đẩy doanh thu bán kèm phụ kiện cao cấp (Upsell) không?"
            },
            {
                "br": "BR-04",
                "scenario": "Tăng số lượng Camera khi đang có Charm Premium: Đang có 1 Cam + 1 Premium (+50k), tăng Cam lên 2",
                "pre": "Giỏ có 1 Cam + 1 Charm Premium (+50k). Khách bấm [+] tăng Camera lên 2.",
                "rule": "Giữ nguyên 1 Charm Premium (+50k) và mở thêm 01 suất mua kèm ưu đãi còn trống để khách tự do chọn thêm món thứ 2. KHÔNG tự ý nhân đôi Premium thành 2 cái (+100k).",
                "alt": "Khách có thể chọn thêm 1 Charm Basic (giá ưu đãi 0đ) hoặc thêm 1 Charm Premium (+50k) theo nhu cầu.",
                "biz_question": "Khi khách tăng số lượng Camera, Kinh doanh muốn hệ thống gợi ý chọn thêm phụ kiện mua kèm ngay hay để khách tự chọn trong giỏ hàng?"
            },
            {
                "br": "BR-05",
                "scenario": "Giảm Camera khi đang có cả 2 loại: Đang có 2 Cam + 1 Basic + 1 Premium (+50k), giảm Cam xuống 1",
                "pre": "Giỏ có 2 Cam + 1 Basic (0đ) + 1 Premium (+50k). Khách bấm [-] giảm Camera về 1.",
                "rule": "Hạn mức ưu đãi giảm còn 1: Ưu tiên giữ lại suất ưu đãi cho Charm Basic (0đ), chiếc Charm Premium chuyển sang tính theo giá bán lẻ gốc 200.000đ.",
                "alt": "Hiển thị thông báo nhanh cho khách chọn: 'Áp dụng ưu đãi cho Charm Basic (0đ)' hoặc 'Áp dụng ưu đãi cho Charm Premium (50k)'.",
                "biz_question": "Kinh doanh muốn ưu tiên giữ suất mua kèm giá 0đ hay giữ suất có phụ thu (+50k) khi khách giảm số lượng mua thiết bị chính?"
            },
            {
                "br": "BR-06",
                "scenario": "Chính sách đổi trả Camera khi có mua kèm phụ kiện phụ thu nâng cấp",
                "pre": "Đơn hàng hoàn tất gồm 1 Camera Play 4 (800k) + 1 Charm Premium mua kèm (50k). Khách muốn trả Camera, giữ lại Charm.",
                "rule": "Khách trả Camera nhưng giữ Charm Premium: Charm không còn điều kiện mua kèm ưu đãi, phải tính theo giá bán lẻ gốc (200k). Số tiền hoàn lại = Tiền Camera (800k) - Chênh lệch giá Charm (200k - 50k = 150k) = 650.000đ.",
                "alt": "Nếu khách hoàn trả cả Camera và Charm Premium còn nguyên vẹn: Hoàn đủ 100% số tiền thực thu (850.000đ).",
                "biz_question": "Kinh doanh có đồng ý với nguyên tắc cấn trừ tiền hoàn này khi khách tách lẻ phụ kiện mua kèm không, hay bắt buộc phải trả trọn bộ mới giải quyết?"
            }
        ]
    },
    {
        "module": "3. Bán chéo số lượng (Mua thêm phụ kiện thứ 2, 3 giảm 50%)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Mua kèm nhiều phụ kiện giá khác nhau: 1 Cam, chọn 3 Charm giá bán lẻ khác nhau (100k, 150k, 200k)",
                "pre": "Chính sách: Mua 1 Cam được 1 Charm giá ưu đãi (hiện tại 0đ) + tối đa 2 Charm tiếp theo giảm 50%. Khách chọn 3 Charm (100k, 150k, 200k).",
                "rule": "Ưu tiên quyền lợi lớn nhất cho khách: Chiếc 200k áp dụng giá mua kèm ưu đãi 0đ (tiết kiệm 200k). Chiếc 150k giảm 50% còn 75k. Chiếc 100k giảm 50% còn 50k. Tổng tiền phụ kiện = 125k (Tiết kiệm 325k).",
                "alt": "Nếu chính sách kinh doanh ưu tiên bảo vệ biên lợi nhuận: Áp dụng giá ưu đãi 0đ cho chiếc rẻ nhất (100k) và giảm 50% cho 2 chiếc còn lại.",
                "biz_question": "Kinh doanh chọn phương án nào: Ưu đãi tốt nhất cho khách (giảm giá món đắt nhất) hay Bảo vệ doanh thu (giảm giá món rẻ nhất)?"
            },
            {
                "br": "BR-02",
                "scenario": "Vượt trần số lượng mua kèm ưu đãi: Mua 1 Cam nhưng khách chọn tới 5 Charm",
                "pre": "1 Camera Play 4. Hạn mức tối đa: 1 Charm giá ưu đãi 0đ + 2 Charm giảm 50% (Tổng tối đa 3 chiếc hưởng giá mua kèm ưu đãi). Khách chọn 5 Charm.",
                "rule": "Chiếc thứ 1: Giá ưu đãi 0đ. Chiếc thứ 2 và 3: Giảm 50%. Chiếc thứ 4 và thứ 5: Tính nguyên giá bán lẻ niêm yết 100%.",
                "alt": "Nếu khách mua thêm Camera thứ 2: Hạn mức mở rộng thành 2 chiếc 0đ + 4 chiếc 50%, chiếc thứ 4 và 5 tự động chuyển sang hưởng giảm 50% ngay lập tức.",
                "biz_question": "Chính sách giảm giá 50% cho phụ kiện thứ 2, thứ 3 có áp dụng xuyên suốt hay chỉ áp dụng trong các đợt chiến dịch khuyến mại nhất định?"
            },
            {
                "br": "BR-03",
                "scenario": "Xóa Camera khi đang có phụ kiện mua kèm giá ưu đãi và giảm 50%",
                "pre": "Giỏ có 1 Camera + 1 Charm giá ưu đãi 0đ + 1 Charm giảm 50%. Khách xóa Camera khỏi giỏ.",
                "rule": "Toàn bộ điều kiện mua kèm bị vô hiệu lực. Hệ thống cảnh báo và chuyển toàn bộ các Charm về giá bán lẻ niêm yết 100% nếu khách muốn giữ giỏ, hoặc xóa toàn bộ phụ kiện.",
                "alt": "Hộp thoại xác nhận: 'Xóa Camera sẽ làm mất toàn bộ giá mua kèm ưu đãi của phụ kiện. Bạn có muốn tiếp tục? [Xóa tất cả] [Giữ Charm giá bán lẻ gốc]'.",
                "biz_question": "Kinh doanh có muốn áp dụng mức chiết khấu thông thường (VD: Giảm 10% phụ kiện lẻ) nếu khách vẫn muốn giữ lại mua phụ kiện sau khi xóa Camera không?"
            },
            {
                "br": "BR-04",
                "scenario": "Dịch chuyển thứ tầng giá ưu đãi khi giảm Camera: Giảm Cam từ 2 xuống 1",
                "pre": "Giỏ có 2 Cam + 2 Charm (ưu đãi 0đ) + 2 Charm (giảm 50%). Khách bấm giảm Camera xuống 1.",
                "rule": "Hạn mức mới: 1 Charm ưu đãi 0đ + 2 Charm giảm 50% + 1 Charm tính giá bán lẻ gốc 100%. Hệ thống tự động dịch chuyển món đắt nhất vào nhóm 0đ, tiếp theo vào nhóm 50% và món còn lại về giá gốc.",
                "alt": "Tổng tiền đơn hàng cập nhật tức thì kèm thông báo giải thích việc điều chỉnh hạn mức mua kèm.",
                "biz_question": "Quy tắc trượt tầng giá mua kèm tự động này có đúng với mong đợi vận hành của đội ngũ Kinh doanh không?"
            },
            {
                "br": "BR-05",
                "scenario": "Đóng gói & Giao nhận khi đơn hàng có nhiều mức giá mua kèm khác nhau",
                "pre": "Đơn hàng hoàn tất gồm 1 Cam + 1 Charm giá ưu đãi 0đ + 2 Charm giảm 50%.",
                "rule": "Tất cả các sản phẩm được đóng chung trong CÙNG MỘT KIỆN HÀNG nếu xuất cùng một kho, không phân biệt phụ kiện mua kèm theo các mức giá khác nhau.",
                "alt": "Hóa đơn và phiếu giao hàng ghi rõ từng dòng sản phẩm với đơn giá mua kèm thực tế để đối soát khi giao nhận.",
                "biz_question": "Kho vận có cần phân biệt tem nhãn giữa phụ kiện mua kèm giá 0đ và phụ kiện mua thêm giảm giá 50% không?"
            }
        ]
    },
    {
        "module": "4. Combo Booster (Mua kèm khi mua gói Combo thiết bị)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Mua Combo đóng gói sẵn: Khách mua gói 'Combo An Ninh: Play 4 + IQ4S'",
                "pre": "Gian hàng mở bán mã Combo đóng gói sẵn có chứa 01 thiết bị Camera Play 4.",
                "rule": "Do trong Combo có chứa thiết bị Camera Play 4, khách vẫn được hưởng đầy đủ 01 suất mua kèm phụ kiện Charm với giá ưu đãi (hiện tại 0đ) như khi mua lẻ Play 4.",
                "alt": "Trang chi tiết của Combo hiển thị rõ: 'Ưu đãi mua kèm 01 Charm Play 4 với giá đặc biệt 0đ'.",
                "biz_question": "Nếu mã Combo này đã được giảm giá sốc hơn 20%, Kinh doanh có muốn loại trừ chính sách mua kèm phụ kiện giá ưu đãi không?"
            },
            {
                "br": "BR-02",
                "scenario": "Gom Combo mềm: Khách tự nhặt 1 Play 4 lẻ và 1 IQ4S lẻ vào giỏ",
                "pre": "Khách chọn riêng lẻ 1 Play 4 và 1 IQ4S vào giỏ hàng.",
                "rule": "Hệ thống tự động nhận diện cặp sản phẩm hợp lệ để áp dụng giá ưu đãi Combo (nếu có), đồng thời vẫn bảo đảm 01 suất mua kèm Charm với giá ưu đãi.",
                "alt": "Hiển thị thông báo: 'Đã tự động ghép Combo An Ninh - Tiết kiệm thêm 100k + Ưu đãi mua kèm 1 Charm giá 0đ'.",
                "biz_question": "Kinh doanh muốn hệ thống tự động ghép Combo kích cầu hay bắt buộc khách phải chọn mua đúng mã Combo được tạo sẵn?"
            },
            {
                "br": "BR-03",
                "scenario": "Phá vỡ Combo mềm: Khách xóa bớt IQ4S, chỉ giữ lại Play 4",
                "pre": "Đang có Combo mềm trong giỏ và đã chọn Charm mua kèm giá ưu đãi. Khách xóa chiếc IQ4S.",
                "rule": "Combo bị hủy, Play 4 quay về giá bán lẻ thông thường. Do Camera Play 4 vẫn có chính sách bán kèm phụ kiện nên suất mua kèm Charm giá ưu đãi VẪN ĐƯỢC GIỮ NGUYÊN.",
                "alt": "Thông báo: 'Combo an ninh đã hủy do giảm bớt sản phẩm. Giá Play 4 quay về giá gốc. Suất mua kèm phụ kiện giá ưu đãi vẫn được bảo lưu'.",
                "biz_question": "Nếu sau này có các gói Combo đặc thù chỉ khi mua Combo mới có suất mua kèm phụ kiện, Kinh doanh có muốn thu hồi quyền mua kèm khi phá Combo không?"
            },
            {
                "br": "BR-04",
                "scenario": "Mua kết hợp cả Combo lẫn mua lẻ: Mua 1 Combo (có 1 Play 4) VÀ thêm 2 Play 4 lẻ",
                "pre": "Giỏ có 1 Combo (chứa 1 Play 4) + 2 Camera Play 4 mua lẻ riêng.",
                "rule": "Cộng dồn hạn mức mua kèm: 1 Play 4 từ Combo + 2 Play 4 lẻ = Tổng cộng 3 Camera Play 4 → Cấp chính xác 3 suất mua kèm Charm theo giá ưu đãi.",
                "alt": "Giao diện hiển thị rõ ràng: 'Bạn có 3 suất mua kèm phụ kiện giá ưu đãi (1 từ Combo An Ninh + 2 từ Camera mua lẻ)'.",
                "biz_question": "Có áp dụng giới hạn số lượng tối đa phụ kiện mua kèm trên một đơn hàng khi khách vừa mua Combo vừa mua lẻ không?"
            }
        ]
    },
    {
        "module": "5. Dịch vụ bám kèm (Mua kèm gắn với gói Cloud)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Điều kiện kép (Lệch số lượng): Mua 3 Camera Play 4 nhưng chỉ mua kèm 1 gói Cloud",
                "pre": "Chính sách quy định điều kiện kép: Mua Camera Play 4 VÀ đăng ký gói Cloud mới được mua kèm Charm giá ưu đãi (0đ).",
                "rule": "Hạn mức mua kèm tính theo số nhỏ nhất (Hàm Min): 3 Camera nhưng chỉ có 1 Cloud → Chỉ có 01 suất mua kèm Charm giá ưu đãi 0đ. 2 Camera còn lại không có Cloud thì không có suất mua kèm giá ưu đãi.",
                "alt": "Hiển thị gợi ý: 'Đăng ký thêm 2 gói Cloud để mở thêm 2 suất mua kèm Charm giá ưu đãi cho 2 Camera còn lại [Chọn gói Cloud]'.",
                "biz_question": "Đối với 2 Camera không kèm Cloud, Kinh doanh có muốn cho phép khách mua Charm theo chính sách giảm giá 50% không?"
            },
            {
                "br": "BR-02",
                "scenario": "Thời hạn gói Cloud: Khách chọn gói Cloud 1 tháng thay vì gói 6 hoặc 12 tháng",
                "pre": "Chính sách quy định chỉ gói Cloud từ 6 tháng trở lên mới đủ điều kiện mua kèm Charm giá ưu đãi 0đ.",
                "rule": "Gói Cloud 1 tháng không đủ điều kiện mua kèm Charm giá ưu đãi. Chỉ khi chọn gói Cloud kỳ hạn >= 6 tháng mới mở suất mua kèm Charm giá ưu đãi (0đ).",
                "alt": "Khi khách chọn gói 1 tháng: Hiển thị tooltip 'Nâng cấp lên gói 6 tháng để nhận ngay quyền mua kèm Charm Play 4 với giá ưu đãi 0đ'.",
                "biz_question": "Kỳ hạn gói Cloud tối thiểu để được hưởng quyền mua kèm Charm giá ưu đãi chốt là 6 tháng hay 12 tháng?"
            },
            {
                "br": "BR-03",
                "scenario": "Hạ gói Cloud trong giỏ: Đang có 1 Cam + Cloud 6T (được giá ưu đãi), khách bấm đổi Cloud xuống 1T",
                "pre": "Đã có 1 Cam + Cloud 6T + Charm giá ưu đãi 0đ. Khách đổi thời hạn gói Cloud từ 6T xuống 1T.",
                "rule": "Mất điều kiện mua kèm ưu đãi: Charm lập tức bị chuyển về giá bán lẻ niêm yết 150.000đ.",
                "alt": "Bảng xác nhận: 'Gói Cloud 1 tháng không đủ điều kiện mua kèm Charm giá ưu đãi. Phụ kiện Charm sẽ được tính theo giá bán lẻ 150.000đ. Bạn có muốn đổi không?'.",
                "biz_question": "Khi khách hạ gói dịch vụ, Kinh doanh muốn chuyển phụ kiện về giá bán lẻ hay tự động hủy chọn phụ kiện khỏi giỏ hàng?"
            },
            {
                "br": "BR-04",
                "scenario": "Hủy dịch vụ Cloud sau thanh toán / Hoàn tiền Cloud có phụ kiện mua kèm",
                "pre": "Khách đã thanh toán đơn gồm Play 4 + Cloud 12T + Charm mua kèm giá ưu đãi 0đ. Sau đó khách yêu cầu hủy gói Cloud.",
                "rule": "Khi khách hủy hợp đồng Cloud trước hạn: Tiền hoàn cước Cloud sẽ bị KHẤU TRỪ chênh lệch giá niêm yết của Charm (150.000đ - 0đ = 150.000đ) do không còn thỏa điều kiện mua kèm ưu đãi.",
                "alt": "Nếu khách không đồng ý khấu trừ tiền: Yêu cầu khách hoàn trả phụ kiện Charm còn nguyên vẹn/nguyên hộp về quầy giao dịch FPT.",
                "biz_question": "Kinh doanh có đồng ý đưa điều khoản ràng buộc khấu trừ phụ kiện mua kèm này vào phụ lục hợp đồng dịch vụ Cloud không?"
            }
        ]
    },
    {
        "module": "6. Mua kèm phiên bản giới hạn theo mốc số lượng (Mốc 3 Camera)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Đạt mốc số lượng: Đang có 2 Camera, khách bấm tăng lên 3 Camera",
                "pre": "Chính sách: Khách mua từ 3 Camera trở lên được mở thêm quyền mua kèm 01 'Charm Bản Giới Hạn / Limited Edition' với mức giá ưu đãi đặc biệt (hiện hành: 0đ).",
                "rule": "Hệ thống tự động mở khóa và thêm 01 Charm Limited Edition (giá ưu đãi 0đ) vào giỏ hàng, đồng thời hạn mức mua kèm Charm tiêu chuẩn tăng lên 3 chiếc.",
                "alt": "Hiển thị thông báo chúc mừng: 'Chúc mừng bạn đã mở khóa quyền mua kèm độc quyền: 01 Charm Limited Edition (0đ)'.",
                "biz_question": "Nếu mẫu Charm Limited Edition bị hết hàng trong kho, Kinh doanh muốn thay thế bằng voucher cước Cloud hay nâng cấp sang phụ kiện khác?"
            },
            {
                "br": "BR-02",
                "scenario": "Rớt mốc số lượng: Đang có 3 Camera + Charm Limited mua kèm, khách giảm Camera xuống 2",
                "pre": "Giỏ có 3 Camera + Charm Limited mua kèm giá ưu đãi 0đ. Khách giảm số lượng Camera về 2.",
                "rule": "Không còn thỏa điều kiện mốc số lượng: Hệ thống tự động thu hồi/xóa Charm Limited khỏi giỏ hàng (vì đây là phiên bản phụ kiện chỉ mở bán kèm khi đạt mốc số lượng).",
                "alt": "Thông báo rõ: 'Số lượng Camera dưới 3, phụ kiện mua kèm Limited Edition đã được tự động gỡ khỏi giỏ hàng'.",
                "biz_question": "Có cần hiển thị cảnh báo trước khi khách bấm giảm số lượng: 'Giảm số lượng sẽ làm mất quyền mua kèm phụ kiện Limited Edition'?"
            },
            {
                "br": "BR-03",
                "scenario": "Hết phụ kiện Limited trong kho ngay lúc khách đang thanh toán",
                "pre": "Khách có 3 Camera + Charm Limited trong giỏ và bấm 'Thanh toán'. Kho vừa hết sạch phiên bản Limited.",
                "rule": "Không chặn đơn hàng Camera chính của khách: Hệ thống xin lỗi và cam kết giao bổ sung phụ kiện Limited sau 7 ngày làm việc hoặc cho phép đổi sang phụ kiện khác tương đương.",
                "alt": "Hộp thoại: 'Phiên bản Charm Limited vừa tạm hết hàng. Chúng tôi cam kết giao bù phụ kiện miễn phí sau 7 ngày. Bạn có đồng ý tiếp tục đặt hàng không?'.",
                "biz_question": "Kinh doanh đồng ý giải pháp cam kết giao bù phụ kiện sau để đảm bảo tỷ lệ chuyển đổi đơn hàng Camera chính hay không?"
            },
            {
                "br": "BR-04",
                "scenario": "Chống đầu cơ gom phụ kiện mua kèm Limited Edition",
                "pre": "Khách hàng hoặc đại lý tạo nhiều đơn hàng liên tục (mỗi đơn 3 Cam) để gom phụ kiện Limited Edition.",
                "rule": "Giới hạn tối đa: Mỗi tài khoản / Số điện thoại / Địa chỉ nhận hàng chỉ được mua kèm tối đa 02 Charm Limited trong suốt toàn bộ chiến dịch.",
                "alt": "Khi đã đạt hạn mức 2 Charm Limited: Các đơn hàng tiếp theo mua 3 Cam vẫn được mua kèm Charm tiêu chuẩn nhưng không mở thêm bản Limited.",
                "biz_question": "Kinh doanh muốn kiểm soát hạn mức này theo Số điện thoại hay Số CCCD/Tài khoản để đảm bảo phụ kiện đến tay người dùng thực tế?"
            }
        ]
    },
    {
        "module": "7. Kế toán & Xuất hóa đơn phụ kiện mua kèm",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Xuất hóa đơn điện tử VAT và ghi nhận chứng từ cho phụ kiện bán kèm",
                "pre": "Đơn hàng thanh toán thành công gồm Camera Play 4 và Charm mua kèm giá ưu đãi (0đ).",
                "rule": "Trên hóa đơn VAT và phiếu bán hàng thể hiện rõ: Charm Play 4 đơn giá niêm yết 150.000đ - Tiền chiết khấu mua kèm: -150.000đ - Thành tiền: 0đ (hoặc theo mức giá mua kèm thực thu) theo đúng quy định đăng ký chương trình xúc tiến thương mại.",
                "alt": "Hệ thống ghi nhận chi phí chênh lệch mua kèm vào ngân sách xúc tiến bán hàng của dự án.",
                "biz_question": "Mẫu thể hiện hóa đơn VAT cho phụ kiện mua kèm giá 0đ đã được phòng Kế toán duyệt theo quy định thuế chưa?"
            },
            {
                "br": "BR-02",
                "scenario": "Áp dụng kèm Mã giảm giá / Voucher toàn sàn (Ví dụ: Giảm 10% tổng đơn)",
                "pre": "Khách nhập mã Voucher 'GIAM10 - Giảm 10% tổng đơn' cho giỏ hàng có Camera và phụ kiện mua kèm.",
                "rule": "Voucher giảm % chỉ áp dụng trên giá trị thực thu của Camera và khoản phụ thu mua kèm (nếu có). Phụ kiện có giá mua kèm 0đ không được tham gia tính giảm thêm.",
                "alt": "Hiển thị chi tiết: 'Voucher GIAM10 áp dụng giảm giá trên Camera Play 4 và phần tiền phụ thu mua kèm'.",
                "biz_question": "Chương trình bán kèm phụ kiện giá ưu đãi có được phép áp dụng chung với các voucher giảm giá/chiết khấu khác trên website không?"
            }
        ]
    },
    {
        "module": "8. Quản trị Tồn kho & Cạnh tranh khi mua kèm phụ kiện",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Hết hàng tức thời: Mẫu Charm mua kèm bị hết kho ngay trước lúc thanh toán",
                "pre": "Khách đang chọn Charm Mẫu A trong giỏ. Trong lúc chờ, khách khác đã mua hết tồn kho Mẫu A.",
                "rule": "Khi bấm 'Thanh toán', nếu mẫu A đã hết hàng, hệ thống nhắc khách đổi sang mẫu Charm khác còn hàng trong hạn mức mua kèm. Không hủy đơn mua Camera của khách.",
                "alt": "Mẫu A bị làm mờ, gắn nhãn 'Tạm hết hàng'. Nổi bật nút 'Chọn mẫu phụ kiện khác thay thế'.",
                "biz_question": "Nếu toàn bộ các mẫu phụ kiện Charm đều hết hàng, Kinh doanh có đồng ý cấp mã voucher để khách mua phụ kiện ở đợt hàng sau không?"
            },
            {
                "br": "BR-02",
                "scenario": "Tồn kho phụ kiện không đủ số lượng mong muốn: Mua 3 Cam muốn 3 Charm A, nhưng kho chỉ còn 1 cái",
                "pre": "Kho còn 1 Charm A và nhiều Charm B, C. Khách muốn chọn cả 3 Charm Mẫu A theo giá mua kèm ưu đãi.",
                "rule": "Giới hạn số lượng chọn tối đa của mỗi mẫu phụ kiện không vượt quá số lượng tồn kho thực tế. Mẫu A chỉ cho chọn tối đa 1 chiếc.",
                "alt": "Thông báo: 'Mẫu A chỉ còn 1 sản phẩm. Vui lòng chọn thêm 2 mẫu phụ kiện khác để tận dụng đủ 3 suất mua kèm ưu đãi'.",
                "biz_question": "Kinh doanh có đồng ý hiển thị số lượng tồn kho còn lại của từng mẫu phụ kiện trên giao diện giỏ hàng không?"
            },
            {
                "br": "BR-03",
                "scenario": "Để giỏ hàng qua đêm (Chính sách giá mua kèm ưu đãi kết thúc hoặc thay đổi)",
                "pre": "Khách bỏ Camera + Charm mua kèm giá 0đ vào giỏ lúc đêm. Sáng hôm sau quay lại thanh toán thì chương trình giá ưu đãi đã kết thúc (chuyển sang giá bán lẻ hoặc giá mua kèm mới).",
                "rule": "Giá bán và quyền lợi mua kèm được tính tại thời điểm bấm Đặt hàng (Point of Sale). Khi chính sách kết thúc, giỏ hàng tự động cập nhật lại theo giá hiện hành.",
                "alt": "Thông báo rõ: 'Chính sách giá mua kèm ưu đãi đã thay đổi. Giỏ hàng của bạn đã được cập nhật lại theo mức giá hiện hành'.",
                "biz_question": "Có chính sách ngoại lệ nào gia hạn mức giá mua kèm cũ thêm một khoảng thời gian cho khách hàng đã lưu giỏ hàng không?"
            }
        ]
    },
    {
        "module": "9. Xung đột Khuyến mãi & Điểm thưởng FPT",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Camera Play 4 tham gia Flash Sale giá sốc",
                "pre": "Camera Play 4 giảm giá đặc biệt từ 800k còn 599k trong khung giờ Flash Sale.",
                "rule": "Theo cấu hình kinh doanh: Nếu Flash Sale đã giảm giá quá sâu, sản phẩm Camera Flash Sale KHÔNG áp dụng kèm quyền mua phụ kiện giá ưu đãi 0đ (phụ kiện tính giá bán lẻ bình thường).",
                "alt": "Trang Flash Sale ghi rõ: 'Giá sốc Flash Sale không áp dụng kèm chính sách mua kèm phụ kiện giá ưu đãi 0đ'.",
                "biz_question": "Kinh doanh quyết định: Khi bán Camera giá Flash Sale có kèm quyền mua phụ kiện giá ưu đãi 0đ không, hay chỉ áp dụng khi bán đúng giá niêm yết?"
            },
            {
                "br": "BR-02",
                "scenario": "Dùng Điểm thưởng FPT Loyalty thanh toán khoản phụ thu mua kèm",
                "pre": "Khách có 50.000 điểm FPT Loyalty, muốn dùng điểm để cấn trừ khoản tiền mua kèm 50k của Charm Premium.",
                "rule": "Cho phép dùng điểm thưởng FPT Loyalty để thanh toán các khoản tiền phụ thu mua kèm hoặc tổng giá trị đơn hàng (1 điểm = 1 VNĐ).",
                "alt": "Mục thanh toán hiển thị: '[x] Dùng 50.000 điểm FPT để thanh toán phụ thu mua kèm (Số tiền cần trả: 0đ)'.",
                "biz_question": "Kinh doanh có đồng ý cho phép khách dùng điểm thưởng FPT Loyalty để thanh toán tiền mua kèm phụ kiện không?"
            }
        ]
    },
    {
        "module": "10. Vận chuyển & Giao nhận Đa kho cho phụ kiện mua kèm",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Lệch kho hàng: Kho gần khách có Camera nhưng hết phụ kiện mua kèm (Phụ kiện ở kho khác)",
                "pre": "Kho gần khách (Hà Nội) có sẵn Camera nhưng hết Charm. Charm đang ở kho Đà Nẵng hoặc HCM.",
                "rule": "Hệ thống tối ưu điều phối: Điều chuyển Charm về để gom giao chung 1 lần; HOẶC Tách làm 2 kiện giao riêng mà KHÔNG thu thêm bất kỳ chi phí vận chuyển nào của khách đối với kiện phụ kiện mua kèm.",
                "alt": "Màn hình đặt hàng thông báo: 'Đơn hàng được chia làm 2 kiện do phụ kiện mua kèm chuyển từ kho khác. Miễn phí vận chuyển cho kiện phụ kiện mua kèm'.",
                "biz_question": "Kinh doanh chọn phương án nào: Tách làm 2 chuyến giao (FPT chịu ship phụ kiện) hay chờ gom đủ hàng mới giao một lần?"
            },
            {
                "br": "BR-02",
                "scenario": "Khách đặt Online và chọn nhận hàng tại Shop FPT gần nhà",
                "pre": "Khách đặt mua Camera kèm phụ kiện trên web và chọn đến Shop FPT nhận hàng.",
                "rule": "Kiểm tra tồn kho tại Shop: Nếu Shop có sẵn Camera nhưng thiếu phụ kiện mua kèm, hệ thống thông báo thời gian điều chuyển phụ kiện về Shop trước khi khách ra lấy.",
                "alt": "Thông báo: 'Cửa hàng có sẵn Camera. Phụ kiện mua kèm dự kiến có mặt tại Shop sau 24h. Chúng tôi sẽ gửi tin nhắn SMS khi trọn bộ sẵn sàng'.",
                "biz_question": "Khách có được phép nhận Camera trước tại quầy, còn phụ kiện mua kèm sẽ gửi chuyển phát về tận nhà miễn phí sau không?"
            }
        ]
    },
    {
        "module": "11. Cổng thanh toán & Trạng thái đơn hàng mua kèm",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Thanh toán Online thất bại hoặc Quá hạn chờ thanh toán (Quá 15 phút)",
                "pre": "Khách chọn thanh toán qua VNPAY/MoMo cho đơn gồm Camera và phụ kiện mua kèm nhưng để quá 15 phút không thanh toán.",
                "rule": "Hạn mức mua kèm ưu đãi và tồn kho phụ kiện được tạm giữ trong 15 phút. Sau 15 phút không hoàn tất, đơn tự hủy và giải phóng tồn kho cho khách khác.",
                "alt": "Giao diện hiển thị: 'Giao dịch đã hết thời gian chờ thanh toán. Giỏ hàng của bạn đã được khôi phục để đặt lại'.",
                "biz_question": "Thời gian giữ chỗ tồn kho và hạn mức mua kèm khi thanh toán trực tuyến nên là 15 phút hay 30 phút?"
            },
            {
                "br": "BR-02",
                "scenario": "Khách hủy đơn hàng COD trước khi xuất kho",
                "pre": "Đơn hàng COD gồm Camera và phụ kiện mua kèm đang ở trạng thái 'Chờ xác nhận'. Khách bấm 'Hủy đơn hàng' trên web.",
                "rule": "Cho phép khách hủy đơn ngay lập tức nếu hàng chưa bàn giao cho shipper. Toàn bộ hạn mức mua kèm ưu đãi và tồn kho phụ kiện được hoàn trả về hệ thống.",
                "alt": "Hiển thị danh sách lý do hủy đơn để khách phản hồi (Đổi ý, chọn nhầm mẫu phụ kiện, muốn đổi địa chỉ).",
                "biz_question": "Sau khi khách tự hủy đơn, có cần gửi tin nhắn kèm ưu đãi giữ chân khách hàng quay lại mua không?"
            }
        ]
    },
    {
        "module": "12. Đổi trả từng phần & Bảo hành phụ kiện mua kèm (RMA & Warranty)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Bảo hành 1 đổi 1 Camera Play 4 do lỗi kỹ thuật của nhà sản xuất",
                "pre": "Camera chính bị lỗi phần cứng trong thời gian bảo hành chính hãng.",
                "rule": "Kỹ thuật viên / Cửa hàng đổi mới 01 Camera Play 4 tương đương cho khách. Khách KHÔNG phải nộp lại phụ kiện Charm đã mua kèm.",
                "alt": "Phiếu bảo hành ghi rõ: 'Đổi mới Camera Play 4 - Khách giữ nguyên phụ kiện mua kèm Charm'.",
                "biz_question": "Nếu bản thân phụ kiện Charm bị lỗi sản xuất (ví dụ gãy khớp nối), Kinh doanh có chính sách bảo hành 1 đổi 1 phụ kiện trong bao nhiêu ngày?"
            },
            {
                "br": "BR-02",
                "scenario": "Khách trả Camera Play 4 để lấy lại tiền mặt (Chính sách trả hàng có phụ kiện mua kèm)",
                "pre": "Đơn hàng gồm 1 Cam (800k) + 1 Charm mua kèm giá ưu đãi 0đ (giá niêm yết 150k). Sau 3 ngày khách yêu cầu trả Camera đòi lại tiền.",
                "rule": "Khách trả Camera chính thì phụ kiện mua kèm không còn thỏa điều kiện hưởng giá ưu đãi: Khách phải hoàn trả lại Camera kèm phụ kiện Charm còn nguyên vẹn. Nếu khách giữ lại Charm hoặc Charm đã hư hỏng, hệ thống sẽ KHẤU TRỪ giá bán lẻ niêm yết của Charm (150.000đ) vào số tiền hoàn lại của Camera (Khách chỉ nhận lại 650.000đ).",
                "alt": "Trường hợp Charm Premium phụ thu 50k: Nếu chỉ trả Camera giữ Charm Premium, hoàn tiền = 800k - (200k - 50k) = 650.000đ.",
                "biz_question": "Kinh doanh chọn phương án xử lý phụ kiện mua kèm khi trả hàng: (A) Khấu trừ theo giá bán lẻ niêm yết; (B) Khấu trừ theo giá vốn; hay (C) Bắt buộc trả đủ trọn bộ mới cho hoàn tiền?"
            }
        ]
    },
    {
        "module": "13. Trải nghiệm Khách vãng lai & Đăng nhập tài khoản khi mua kèm",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Khách vãng lai chọn Camera và phụ kiện mua kèm rồi mới bấm Đăng nhập",
                "pre": "Khách chưa đăng nhập chọn 1 Cam + 1 Charm A (giá mua kèm ưu đãi). Sau đó bấm Đăng nhập. Tài khoản đã có sẵn 1 Cam từ trước.",
                "rule": "Hợp nhất giỏ hàng thông minh: Tổng số Camera = 1 (vãng lai) + 1 (tài khoản) = 2 Camera. Hạn mức mua kèm phụ kiện giá ưu đãi tự động nâng lên 2. Giữ nguyên Charm A và mở thêm 1 suất mua kèm ưu đãi thứ 2.",
                "alt": "Thông báo: 'Giỏ hàng đã được đồng bộ. Bạn hiện có 2 Camera và 2 suất mua kèm phụ kiện giá ưu đãi! [Chọn thêm phụ kiện]'.",
                "biz_question": "Nếu giỏ hàng cũ trên tài khoản đã có sẵn Charm B, Kinh doanh muốn hệ thống tự gộp thành (Charm A + Charm B = 2 phụ kiện mua kèm ưu đãi) luôn không?"
            },
            {
                "br": "BR-02",
                "scenario": "Khách đăng xuất tài khoản khi đang có giỏ hàng phụ kiện mua kèm",
                "pre": "Khách đang có giỏ hàng Camera + phụ kiện mua kèm giá ưu đãi, sau đó bấm 'Đăng xuất'.",
                "rule": "Giỏ hàng cùng các điều kiện mua kèm ưu đãi của tài khoản được lưu an toàn trên hệ thống. Trình duyệt công cộng được làm sạch giỏ hàng để bảo mật.",
                "alt": "Khi đăng nhập lại từ bất kỳ thiết bị nào, giỏ hàng cùng các quyền lợi mua kèm ưu đãi được phục hồi nguyên vẹn.",
                "biz_question": "Giỏ hàng mua kèm phụ kiện lưu trên tài khoản khách hàng nên được lưu giữ tối đa bao nhiêu ngày trước khi tự động làm mới?"
            }
        ]
    },
    {
        "module": "14. Mua Nhanh (Buy Now) & Khối Bán Kèm trên Trang Sản Phẩm (PDP)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Khách bấm nút 'MUA NGAY' (Express Checkout) trên trang Camera Play 4",
                "pre": "Khách đang xem trang chi tiết Camera Play 4 và bấm nút 'MUA NGAY' để đi thẳng đến bước thanh toán.",
                "rule": "Bật Pop-up nhanh (Quick Add-on Modal) trước khi chuyển hướng: 'Bạn được tặng 01 suất mua kèm Charm ưu đãi 0đ. Bạn có muốn nhận ngay không?' kèm danh sách mẫu Charm còn hàng và 2 nút: [Nhận phụ kiện 0đ & Thanh toán] hoặc [Bỏ qua & Đến thanh toán].",
                "alt": "Nếu khách bấm 'Bỏ qua': Hệ thống chuyển thẳng sang Checkout chỉ với Camera Play 4, không ép buộc chọn phụ kiện.",
                "biz_question": "Kinh doanh muốn bật Pop-up chọn nhanh phụ kiện khi bấm 'Mua ngay' hay chuyển thẳng sang Checkout rồi nhắc chọn phụ kiện tại trang Checkout?"
            },
            {
                "br": "BR-02",
                "scenario": "Khối hiển thị mua kèm (Attached Deal Widget) trên trang chi tiết sản phẩm",
                "pre": "Khách xem Camera Play 4, bên dưới có khối 'Ưu đãi mua kèm phụ kiện độc quyền'.",
                "rule": "Mặc định KHÔNG tick sẵn phụ kiện để tránh khách không có nhu cầu hoặc chọn nhầm màu. Khách tick chọn mẫu Charm mong muốn thì nút 'Thêm vào giỏ' sẽ tự động gom cả Camera và Charm vào giỏ hàng trong 1 lần bấm.",
                "alt": "Nếu khách tăng số lượng Camera lên 2 tại PDP: Hiển thị tooltip 'Bạn được mua kèm 2 Charm giá ưu đãi! Bạn có thể chọn thêm mẫu khác trong giỏ hàng'.",
                "biz_question": "Kinh doanh muốn mặc định tick sẵn 1 mẫu phụ kiện hay để trống cho khách tự chủ động chọn?"
            },
            {
                "br": "BR-03",
                "scenario": "Biến thể phụ kiện mua kèm bị Hết hàng (Out of Stock) tại trang sản phẩm",
                "pre": "Charm Mẫu A (Màu Hồng) hết hàng trong kho, các mẫu B và C vẫn còn hàng.",
                "rule": "Tại khối mua kèm trên PDP: Mẫu A hiển thị dưới dạng làm mờ (Grey-out), gắn nhãn 'Tạm hết hàng' và khóa không cho tick chọn. Các mẫu còn hàng hiển thị bình thường kèm nhãn 'Giá ưu đãi 0đ'.",
                "alt": "Nếu tất cả các mẫu Charm đều hết hàng: Khối mua kèm tự động ẩn khỏi trang PDP để tránh gây bức xúc cho khách hàng.",
                "biz_question": "Khi một mẫu phụ kiện hết hàng, Kinh doanh muốn làm mờ gắn nhãn 'Tạm hết' hay ẩn hoàn toàn mẫu đó khỏi trang sản phẩm?"
            }
        ]
    },
    {
        "module": "15. Thao tác Bán kèm trên Giỏ Hàng Trượt (Mini-Cart / Drawer Cart)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Đổi mẫu phụ kiện mua kèm trực tiếp trên ngăn kéo Mini-Cart",
                "pre": "Khách thêm Camera + 1 Charm A vào giỏ. Mini-cart bên phải màn hình trượt ra.",
                "rule": "Tại dòng Charm A trên Mini-cart, hiển thị nút bấm nhỏ '[Đổi mẫu]'. Khi khách bấm, mở danh sách popover các mẫu Charm còn hàng để khách chuyển đổi sang Charm B ngay trên Mini-cart mà không cần chuyển trang.",
                "alt": "Nếu khách chọn đổi sang Charm Premium (+50k): Subtotal của Mini-cart tự động cập nhật cộng thêm 50.000đ ngay lập tức.",
                "biz_question": "Kinh doanh có đồng ý cho khách đổi mẫu phụ kiện ngay trên Mini-cart không hay bắt buộc phải vào trang Giỏ hàng chi tiết (/cart)?"
            },
            {
                "br": "BR-02",
                "scenario": "Tăng số lượng Camera trực tiếp trên Mini-Cart",
                "pre": "Mini-cart đang có 1 Camera + 1 Charm A (0đ). Khách bấm nút [+] tăng Camera lên 2.",
                "rule": "Hạn mức ưu đãi tăng lên 2. Giữ nguyên 1 Charm A đã chọn, đồng thời bên dưới xuất hiện 1 dòng slot trống nổi bật: '+ Thêm 1 Charm ưu đãi 0đ [Chọn ngay]'.",
                "alt": "Nếu khách bấm 'Thanh toán' mà chưa chọn slot thứ 2: Đơn hàng vẫn tiếp tục bình thường với 2 Camera + 1 Charm A.",
                "biz_question": "Khi tăng Camera trên Mini-cart, Kinh doanh muốn hệ thống tự động nhân đôi mẫu Charm hiện tại hay mở slot trống để khách tự chọn?"
            }
        ]
    },
    {
        "module": "16. Ràng buộc Khuyến mãi & Kho tại Trang Thanh Toán (Checkout)",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Lời nhắc cuối (Last Reminder) tại Trang Thanh toán khi chưa chọn phụ kiện",
                "pre": "Khách mua Camera Play 4 nhưng chưa chọn phụ kiện mua kèm nào và đã đến trang Checkout điền địa chỉ.",
                "rule": "Ngay phía trên nút 'Đặt hàng', hiển thị Banner ưu đãi 1 chạm: 'Bạn đang bỏ lỡ 01 Charm Play 4 giá ưu đãi 0đ! [Nhận Charm ngay]'. Khách click chọn mẫu là hệ thống tự add vào đơn ngay tại Checkout mà không reload trang.",
                "alt": "Khách có quyền bấm 'Bỏ qua' hoặc nhấn thẳng nút 'Đặt hàng' nếu không có nhu cầu.",
                "biz_question": "Có cần hiển thị Banner nhắc nhận phụ kiện 0đ ngay tại bước Thanh toán cuối cùng để tối ưu tỷ lệ gắn kèm sản phẩm không?"
            },
            {
                "br": "BR-02",
                "scenario": "Khách nhập Mã giảm giá (Coupon Code) xung đột với suất mua kèm",
                "pre": "Khách có Camera và Charm (0đ) trong đơn, khách nhập mã voucher toàn sàn (VD: Mã GIAM50K cho đơn từ 1.000.000đ).",
                "rule": "Giá trị của Charm 0đ KHÔNG được tính vào điều kiện giá trị đơn hàng tối thiểu (Min Basket Size) để kích hoạt voucher. Voucher chỉ giảm trên số tiền thực thu của Camera và phần phụ thu mua kèm (nếu có).",
                "alt": "Nếu khách nhập voucher chuyên biệt dành cho phụ kiện: Hệ thống áp dụng quy tắc ưu đãi tốt nhất, không cộng dồn hai mức giảm giá cho cùng 1 món.",
                "biz_question": "Phụ kiện mua kèm giá 0đ có được tính vào tổng giá trị đơn hàng để xét điều kiện áp dụng mã voucher toàn sàn không?"
            },
            {
                "br": "BR-03",
                "scenario": "Khách đổi Địa chỉ nhận hàng tại Checkout dẫn đến Lệch kho tồn phụ kiện",
                "pre": "Khách đổi địa chỉ giao hàng từ Hà Nội (kho còn Charm) sang tỉnh khác (kho tỉnh hết Charm, chỉ còn kho trung tâm).",
                "rule": "Không tự ý gỡ Charm của khách: Hệ thống hiển thị thông báo thời gian giao hàng: 'Phụ kiện mua kèm sẽ được điều chuyển từ kho trung tâm, đơn hàng dự kiến giao sau 2-3 ngày làm việc (Miễn phí vận chuyển)'.",
                "alt": "Nếu khách cần nhận gấp Camera trong 2h: Cho phép khách tick chọn 'Tách giao Camera trước, phụ kiện giao sau'.",
                "biz_question": "Khi kho giao hàng tại địa chỉ của khách hết phụ kiện mua kèm, Kinh doanh muốn điều chuyển kho miễn phí hay thông báo khách đổi mẫu khác?"
            }
        ]
    },
    {
        "module": "17. Bán kèm Sau Đặt Hàng (Post-Purchase One-Click Upsell) tại Trang Thành Công",
        "cases": [
            {
                "br": "BR-01",
                "scenario": "Khách hoàn tất thanh toán Camera mà quên chọn phụ kiện mua kèm",
                "pre": "Khách đã bấm Đặt hàng thành công đơn Camera Play 4 và chuyển đến trang 'Cảm ơn / Đặt hàng thành công'.",
                "rule": "Kích hoạt đồng hồ đếm ngược 5 phút tại trang Cảm ơn: 'Đặc quyền của bạn: Thêm 01 Charm Play 4 với giá ưu đãi 0đ vào đơn hàng vừa tạo (Miễn phí vận chuyển)'. Khách click [Thêm ngay], hệ thống tự động gộp vào đơn hàng vừa tạo trước khi đẩy sang kho WMS.",
                "alt": "Sau 5 phút hoặc khi đơn hàng đã chuyển trạng thái 'Đang đóng gói', ưu đãi bổ sung này tự động hết hiệu lực.",
                "biz_question": "Kinh doanh có đồng ý cho phép khách bổ sung phụ kiện mua kèm vào đơn hàng vừa tạo trong vòng 5 phút tại màn hình Cảm ơn không?"
            }
        ]
    }
]

xlsx_path = "/Users/hathuy/Documents/FPT-1/thuyttdoc/Kich_Ban_Ky_Thuat_Play4_Charm.xlsx"
tsv_path = "/Users/hathuy/Documents/FPT-1/thuyttdoc/Kich_Ban_Ky_Thuat_Play4_Charm.tsv"
csv_path = "/Users/hathuy/Documents/FPT-1/thuyttdoc/Kich_Ban_Ky_Thuat_Play4_Charm_utf8.csv"
csv_std_path = "/Users/hathuy/Documents/FPT-1/thuyttdoc/Kich_Ban_Ky_Thuat_Play4_Charm.csv"

wb = openpyxl.Workbook()
ws_default = wb.active
wb.remove(ws_default)

ws = wb.create_sheet(title="Kich Ban Chuan - Tech Spec")

font_title = Font(name="Times New Roman", size=16, bold=True, color="1B365D")
font_subtitle = Font(name="Times New Roman", size=12, italic=True, color="555555")
font_header = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
font_section = Font(name="Times New Roman", size=13, bold=True, color="1B365D")
font_body = Font(name="Times New Roman", size=11, color="000000")
font_br = Font(name="Times New Roman", size=11, bold=True, color="004085")
font_q = Font(name="Times New Roman", size=11, italic=True, color="856404")

fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
fill_section = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
fill_zebra = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
fill_q = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

thin_side = Side(style="thin", color="D0D7DE")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

ws.cell(1, 1, "BẢNG ĐẶC TẢ KỊCH BẢN GIỎ HÀNG: MUA KÈM PHỤ KIỆN (ATTACHED ADD-ON)")
ws.cell(1, 1).font = font_title
ws.cell(2, 1, "Dự án: FPT Camera Play 4 & Phụ kiện Charm Add-on Engine | Bản chất: Mua kèm có điều kiện với giá ưu đãi (cấu hình linh hoạt)")
ws.cell(2, 1).font = font_subtitle

headers = [
    "STT",
    "Mã Rule",
    "Nhóm Kịch Bản",
    "Tình Huống Mua Hàng (Customer Scenario)",
    "Điều Kiện Áp Dụng (Pre-conditions)",
    "Quy Tắc Nghiệp Vụ & Quyền Lợi KH (Business Logic & Pricing)",
    "Xử Lý Khi Có Thay Đổi (Alternative & Exception Flow)",
    "Câu Hỏi Khai Vấn & Mong Muốn Đội Kinh Doanh (Questions for Sales/Biz Team)"
]

header_row = 4
for col_idx, h in enumerate(headers, 1):
    c = ws.cell(header_row, col_idx, h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border_cell
ws.row_dimensions[header_row].height = 42

current_row = 5
stt_counter = 1

tsv_rows = []
csv_rows = []

tsv_rows.append(headers)
csv_rows.append(headers)

for mod_idx, mod in enumerate(modules_biz, 1):
    sec_cell = ws.cell(current_row, 1, f"PHẦN {mod_idx}: {mod['module'].upper()}")
    sec_cell.font = font_section
    sec_cell.fill = fill_section
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
    for c_i in range(1, len(headers) + 1):
        ws.cell(current_row, c_i).border = border_cell
        ws.cell(current_row, c_i).fill = fill_section
    ws.row_dimensions[current_row].height = 28
    current_row += 1
    
    for case in mod["cases"]:
        row_data = [
            stt_counter,
            case["br"],
            mod["module"],
            case["scenario"],
            case["pre"],
            case["rule"],
            case["alt"],
            case["biz_question"]
        ]
        
        tsv_rows.append(row_data)
        csv_rows.append(row_data)
        
        is_even = (stt_counter % 2 == 0)
        
        for col_idx, val in enumerate(row_data, 1):
            c = ws.cell(current_row, col_idx, val)
            c.font = font_body
            c.border = border_cell
            
            if is_even:
                c.fill = fill_zebra
                
            if col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx == 2:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = font_br
            elif col_idx == 8:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c.font = font_q
                if not is_even:
                    c.fill = fill_q
            else:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
        ws.row_dimensions[current_row].height = 70
        current_row += 1
        stt_counter += 1

col_widths = {
    "A": 7,    # STT
    "B": 11,   # Mã Rule
    "C": 22,   # Nhóm Kịch Bản
    "D": 32,   # Tình Huống Mua Hàng
    "E": 28,   # Điều Kiện
    "F": 42,   # Quy Tắc Nghiệp Vụ
    "G": 34,   # Xử Lý Khi Có Thay Đổi
    "H": 38    # Câu Hỏi Khai Vấn Kinh Doanh
}

for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

wb.save(xlsx_path)
print(f"Exported Excel to {xlsx_path}")

with open(tsv_path, "w", encoding="utf-8") as f:
    for r in tsv_rows:
        cleaned = [str(cell).replace("\t", " ").replace("\r\n", " ").replace("\n", " ") for cell in r]
        f.write("\t".join(cleaned) + "\n")
print(f"Exported TSV to {tsv_path}")

with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    for r in csv_rows:
        writer.writerow(r)
print(f"Exported CSV UTF8 to {csv_path}")

with open(csv_std_path, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    for r in csv_rows:
        writer.writerow(r)
print(f"Exported CSV to {csv_std_path}")
